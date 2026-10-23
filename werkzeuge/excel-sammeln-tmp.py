# Einmalig: je Typ zusammenfassen, wo welche Platz-Nummer und welche
# Beschriftung im Excel steht (ohne Inhalte der Teile).
import re, sys, os, collections, zipfile
import openpyxl, xlrd
from openpyxl.utils import get_column_letter
TNR = re.compile(r"^T\d{1,4}$")
def label_ok(t):
    return (t.endswith(":") and len(t) <= 40) or (not re.search(r"\d", t) and len(t) <= 30)
je = collections.defaultdict(lambda: {"n": 0, "t": collections.defaultdict(collections.Counter), "farbe": collections.defaultdict(collections.Counter),
     "lab": collections.defaultdict(collections.Counter), "titel": collections.Counter(), "hd": 0, "bilder": 0, "fehler": 0, "max": collections.Counter()})
for zeile in open(sys.argv[1]):
    typ, datei = zeile.split()
    p = "x/" + datei
    d = je[typ]
    try:
        if datei.endswith(".xlsx"):
            wb = openpyxl.load_workbook(p)
            ws = next(w for w in wb.worksheets if w.sheet_state == "visible")
            zellen = []
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None: continue
                    f = c.fill.fgColor.rgb if c.fill and c.fill.fill_type == "solid" and isinstance(c.fill.fgColor.rgb, str) else ""
                    zellen.append((c.coordinate, str(c.value).strip(), f[-6:]))
            with zipfile.ZipFile(p) as z:
                for n in z.namelist():
                    if n.startswith("xl/drawings/drawing") and n.endswith(".xml"):
                        x = z.read(n).decode("utf8", "ignore")
                        d["hd"] += len(re.findall(r"<a:t>\s*HD\s*</a:t>", x)); d["bilder"] += x.count("<xdr:pic>")
        else:
            wb = xlrd.open_workbook(p)
            ws = wb.sheet_by_index(0)
            zellen = [(f"{xlrd.colname(s)}{r+1}", str(ws.cell_value(r, s)).strip(), "") for r in range(ws.nrows) for s in range(ws.ncols) if str(ws.cell_value(r, s)).strip()]
        d["n"] += 1
        maxz = 0
        for ort, t, f in zellen:
            if t.startswith("="): continue
            z = int(re.sub(r"\D", "", ort)); maxz = max(maxz, z)
            if TNR.match(t):
                d["t"][t][ort] += 1
                if f: d["farbe"][t][f] += 1
            elif ort == "A1" or ort == "B1":
                d["titel"][re.sub(r"\d", "9", t)] += 1
            elif label_ok(t) and (ort[0] in "AB" or t.endswith(":")):
                d["lab"][t][ort] += 1
        d["max"][maxz] += 1
    except Exception as e:
        d["fehler"] += 1; print("FEHLER", typ, datei, e)
def nr(t): return int(t[1:])
for typ, d in sorted(je.items()):
    print("#####", typ, "Dateien", d["n"], "Fehler", d["fehler"], "HD-Felder", d["hd"], "Bilder", d["bilder"], "Zeilen", dict(d["max"].most_common(4)))
    print("  Titel:", d["titel"].most_common(4))
    for t in sorted(d["t"], key=nr):
        print("  ", t, sum(d["t"][t].values()), d["t"][t].most_common(3), dict(d["farbe"][t].most_common(2)))
    for l, c in sorted(d["lab"].items(), key=lambda x: -sum(x[1].values())):
        if sum(c.values()) >= 2: print("   L", repr(l), sum(c.values()), c.most_common(3))
